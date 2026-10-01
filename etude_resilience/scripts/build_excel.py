"""Construit la base Excel (données brutes + formules). À recalculer ensuite avec LibreOffice."""
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.table import Table, TableStyleInfo

sys.path.insert(0, str(Path(__file__).parent))
import data as D  # noqa: E402

OUT = Path(__file__).resolve().parents[1] / "Base_resilience_filieres_CI.xlsx"

HEAD = PatternFill("solid", start_color="1F4E3D")
INPUT_FONT = Font(color="0000FF")
HFONT = Font(bold=True, color="FFFFFF")
thin = Side(style="thin", color="BFBFBF")


def header(ws, row, cols, widths=None):
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=j, value=c)
        cell.fill, cell.font = HEAD, HFONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    if widths:
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[L(j)].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def col_index(cols):
    return {c: L(i + 1) for i, c in enumerate(cols)}


wb = Workbook()

# ---------------------------------------------------------------- LISEZMOI
ws = wb.active
ws.title = "LISEZMOI"
lines = [
    "BASE DE DONNÉES – RÉSILIENCE DES FILIÈRES AGRICOLES IVOIRIENNES AUX CHOCS DE PRIX INTERNATIONAUX",
    "MINADERPV – Cellule d'études – version du 1er octobre 2026",
    "",
    "Onglets :",
    "SERIES : séries annuelles par filière/pays (données brutes en bleu, calculs en noir).",
    "EPISODES : épisodes de choc T0 -> T2 ; transmission, amortissement, recette, délais.",
    "INDICATEURS : batterie harmonisée (volatilité, beta baisse/hausse, drawdown, décomposition prix/volume de la recette).",
    "STRESS_PARAM / STRESS_TEST : test de stress standardisé (5 scénarios x 6 filières).",
    "MECANISMES : benchmark international des instruments.",
    "LITTERATURE : revue scientifique (tableau interne).",
    "CARTE_QUALI : filières non quantifiées.",
    "",
    "Conventions :",
    "Transmission = variation prix producteur / variation prix international (même monnaie locale, même période).",
    "Amortissement = 1 - transmission. Recette = prix producteur x quantité (Mds de monnaie locale).",
    "Prix internationaux convertis en monnaie locale/kg : prix_int x fx x conv.",
    "Qualité : A = vérifié sur source officielle/presse 2026 ; B = valeur publiée à revérifier ; C = estimation à valider.",
    "Toutes les figures de la note ministérielle sont générées à partir de ce classeur (scripts/build_charts.py).",
    "Voir AUDIT_DONNEES.md pour les limites et les points à valider.",
]
for i, t in enumerate(lines, 1):
    ws.cell(row=i, column=1, value=t).font = Font(bold=(i == 1), size=12 if i == 1 else 10)
ws.column_dimensions["A"].width = 130

# ---------------------------------------------------------------- SERIES
ws = wb.create_sheet("SERIES")
calc_cols = ["prix_int_local_kg", "part_producteur", "recette_mds", "var_int", "var_prod", "var_recette",
             "transmission_annuelle", "max_prod_cumul", "drawdown_prod", "dlnP", "dlnQ", "dlnR"]
cols = D.SERIES_COLS + calc_cols
C = col_index(cols)
header(ws, 1, cols, [12, 14, 10, 7, 9, 8, 8, 8, 10, 7, 9, 40, 18, 30, 30, 40] + [12] * len(calc_cols))
groups = {}
for i, r in enumerate(D.SERIES, 2):
    for k in D.SERIES_COLS:
        c = ws[f"{C[k]}{i}"]
        c.value = r[k]
        if k in ("prix_int", "fx", "conv", "prix_prod", "prod_kt"):
            c.font = INPUT_FONT
    key = (r["filiere"], r["pays"])
    first = key not in groups
    groups.setdefault(key, [i, i])[1] = i
    g0 = groups[key][0]
    p = i - 1
    ws[f"{C['prix_int_local_kg']}{i}"] = f"={C['prix_int']}{i}*{C['fx']}{i}*{C['conv']}{i}"
    ws[f"{C['part_producteur']}{i}"] = f"={C['prix_prod']}{i}/{C['prix_int_local_kg']}{i}"
    ws[f"{C['recette_mds']}{i}"] = f"={C['prix_prod']}{i}*{C['prod_kt']}{i}/1000"
    if not first:
        ws[f"{C['var_int']}{i}"] = f"={C['prix_int_local_kg']}{i}/{C['prix_int_local_kg']}{p}-1"
        ws[f"{C['var_prod']}{i}"] = f"={C['prix_prod']}{i}/{C['prix_prod']}{p}-1"
        ws[f"{C['var_recette']}{i}"] = f"={C['recette_mds']}{i}/{C['recette_mds']}{p}-1"
        ws[f"{C['transmission_annuelle']}{i}"] = (
            f'=IF(ABS({C["var_int"]}{i})>=0.05,{C["var_prod"]}{i}/{C["var_int"]}{i},"")')
        ws[f"{C['dlnP']}{i}"] = f"=LN({C['prix_prod']}{i}/{C['prix_prod']}{p})"
        ws[f"{C['dlnQ']}{i}"] = f"=LN({C['prod_kt']}{i}/{C['prod_kt']}{p})"
        ws[f"{C['dlnR']}{i}"] = f"=LN({C['recette_mds']}{i}/{C['recette_mds']}{p})"
    ws[f"{C['max_prod_cumul']}{i}"] = f"=MAX({C['prix_prod']}${g0}:{C['prix_prod']}{i})"
    ws[f"{C['drawdown_prod']}{i}"] = f"={C['prix_prod']}{i}/{C['max_prod_cumul']}{i}-1"
    for k in ("part_producteur", "var_int", "var_prod", "var_recette", "drawdown_prod"):
        ws[f"{C[k]}{i}"].number_format = "0.0%"
    for k in ("prix_int_local_kg", "recette_mds", "transmission_annuelle"):
        ws[f"{C[k]}{i}"].number_format = "#,##0.00"
SERIES_C, SERIES_GROUPS = C, groups
n_series = len(D.SERIES) + 1
ws.add_table(Table(displayName="T_SERIES", ref=f"A1:{L(len(cols))}{n_series}",
                   tableStyleInfo=TableStyleInfo(name="TableStyleLight9", showRowStripes=True)))

# ---------------------------------------------------------------- EPISODES
ws = wb.create_sheet("EPISODES")
ecalc = ["int_local_t0", "int_local_t2", "choc_int_usd", "choc_int", "choc_prod", "transmission",
         "amortissement", "recette_t0", "recette_t2", "var_recette", "part_prod_t0", "part_prod_t2"]
ecols = D.EPISODE_COLS + ecalc
E = col_index(ecols)
header(ws, 1, ecols, [11, 12, 13, 7, 22, 9, 9, 8, 8, 7, 7, 7, 7, 8, 8, 8, 7, 7, 8, 8, 24, 40, 9, 30, 10, 26, 30, 40] + [11] * len(ecalc))
for i, r in enumerate(D.EPISODES, 2):
    for k in D.EPISODE_COLS:
        c = ws[f"{E[k]}{i}"]
        c.value = r[k]
        if k in ("int_t0", "int_t2", "fx_t0", "fx_t2", "prod_t0", "prod_t2", "q_t0", "q_t2", "cout_mds_fcfa",
                 "delai_reaction_mois", "delai_recup_mois"):
            c.font = INPUT_FONT
    ws[f"{E['int_local_t0']}{i}"] = f"={E['int_t0']}{i}*{E['fx_t0']}{i}*{E['conv']}{i}"
    ws[f"{E['int_local_t2']}{i}"] = f"={E['int_t2']}{i}*{E['fx_t2']}{i}*{E['conv']}{i}"
    ws[f"{E['choc_int_usd']}{i}"] = f"={E['int_t2']}{i}/{E['int_t0']}{i}-1"
    ws[f"{E['choc_int']}{i}"] = f"={E['int_local_t2']}{i}/{E['int_local_t0']}{i}-1"
    ws[f"{E['choc_prod']}{i}"] = f"={E['prod_t2']}{i}/{E['prod_t0']}{i}-1"
    ws[f"{E['transmission']}{i}"] = f"={E['choc_prod']}{i}/{E['choc_int']}{i}"
    ws[f"{E['amortissement']}{i}"] = f"=1-{E['transmission']}{i}"
    ws[f"{E['recette_t0']}{i}"] = f"={E['prod_t0']}{i}*{E['q_t0']}{i}/1000"
    ws[f"{E['recette_t2']}{i}"] = f"={E['prod_t2']}{i}*{E['q_t2']}{i}/1000"
    ws[f"{E['var_recette']}{i}"] = f"={E['recette_t2']}{i}/{E['recette_t0']}{i}-1"
    ws[f"{E['part_prod_t0']}{i}"] = f"={E['prod_t0']}{i}/{E['int_local_t0']}{i}"
    ws[f"{E['part_prod_t2']}{i}"] = f"={E['prod_t2']}{i}/{E['int_local_t2']}{i}"
    for k in ("choc_int_usd", "choc_int", "choc_prod", "var_recette", "part_prod_t0", "part_prod_t2"):
        ws[f"{E[k]}{i}"].number_format = "0.0%"
    for k in ("transmission", "amortissement"):
        ws[f"{E[k]}{i}"].number_format = "0.00"
n_ep = len(D.EPISODES) + 1
ws.add_table(Table(displayName="T_EPISODES", ref=f"A1:{L(len(ecols))}{n_ep}",
                   tableStyleInfo=TableStyleInfo(name="TableStyleLight9", showRowStripes=True)))

# ---------------------------------------------------------------- INDICATEURS
ws = wb.create_sheet("INDICATEURS")
icols = ["filiere", "pays", "periode", "n_obs", "I1_CV_prix_int", "I2_CV_prix_prod", "CV_recette",
         "I3_beta_baisse_annuel", "I4_beta_hausse_annuel", "I3b_beta_baisse_episodes", "I4b_beta_hausse_episodes",
         "I5_amortissement_baisses", "I6_drawdown_max_prod", "I7_baisse_max_recette", "part_variance_recette_due_Q",
         "correl_dlnP_dlnQ", "I8_recup_max_mois", "I9_cout_observe_mds", "nb_chocs_int_15pct", "nb_chocs_int_20pct"]
header(ws, 1, icols, [12, 14, 12] + [13] * (len(icols) - 3))
S, ep_last = SERIES_C, n_ep
rng = lambda col, a, b: f"SERIES!{S[col]}{a}:{S[col]}{b}"  # noqa: E731
erng = lambda col: f"EPISODES!${E[col]}$2:${E[col]}${ep_last}"  # noqa: E731
for i, ((fil, pays), (a, b)) in enumerate(SERIES_GROUPS.items(), 2):
    ws[f"A{i}"], ws[f"B{i}"] = fil, pays
    ws[f"C{i}"] = f'={rng("campagne", a, a)}&" – "&{rng("campagne", b, b)}'
    ws[f"D{i}"] = b - a + 1
    ws[f"E{i}"] = f"=STDEV({rng('prix_int_local_kg', a, b)})/AVERAGE({rng('prix_int_local_kg', a, b)})"
    ws[f"F{i}"] = f"=STDEV({rng('prix_prod', a, b)})/AVERAGE({rng('prix_prod', a, b)})"
    ws[f"G{i}"] = f"=STDEV({rng('recette_mds', a, b)})/AVERAGE({rng('recette_mds', a, b)})"
    ws[f"H{i}"] = (f'=IFERROR(AVERAGEIFS({rng("transmission_annuelle", a + 1, b)},{rng("var_int", a + 1, b)},"<-0.05"),"")')
    ws[f"I{i}"] = (f'=IFERROR(AVERAGEIFS({rng("transmission_annuelle", a + 1, b)},{rng("var_int", a + 1, b)},">0.05"),"")')
    ws[f"J{i}"] = (f'=IFERROR(AVERAGEIFS({erng("transmission")},{erng("filiere")},A{i},{erng("pays")},B{i},{erng("sens")},"baisse"),"")')
    ws[f"K{i}"] = (f'=IFERROR(AVERAGEIFS({erng("transmission")},{erng("filiere")},A{i},{erng("pays")},B{i},{erng("sens")},"hausse"),"")')
    ws[f"L{i}"] = f'=IF(ISNUMBER(H{i}),1-H{i},"")'
    ws[f"M{i}"] = f"=MIN({rng('drawdown_prod', a, b)})"
    ws[f"N{i}"] = f"=MIN({rng('var_recette', a + 1, b)})"
    ws[f"O{i}"] = (f"=VAR({rng('dlnQ', a + 1, b)})/(VAR({rng('dlnP', a + 1, b)})+VAR({rng('dlnQ', a + 1, b)}))")
    ws[f"P{i}"] = f"=CORREL({rng('dlnP', a + 1, b)},{rng('dlnQ', a + 1, b)})"
    ws[f"Q{i}"] = (f'=IFERROR(_xlfn.MAXIFS({erng("delai_recup_mois")},{erng("filiere")},A{i},{erng("pays")},B{i}),"")')
    ws[f"R{i}"] = (f'=SUMIFS({erng("cout_mds_fcfa")},{erng("filiere")},A{i},{erng("pays")},B{i})')
    ws[f"S{i}"] = f'=COUNTIF({rng("var_int", a + 1, b)},"<=-0.15")'
    ws[f"T{i}"] = f'=COUNTIF({rng("var_int", a + 1, b)},"<=-0.2")'
    for c in "EFGLMNO":
        ws[f"{c}{i}"].number_format = "0.0%"
    for c in "HIJKP":
        ws[f"{c}{i}"].number_format = "0.00"
n_ind = len(SERIES_GROUPS) + 1
# Indicateurs issus des seuls épisodes (filières sans série annuelle complète)
r0 = n_ind + 3
ws[f"A{r0 - 1}"] = "Transmission par épisode (toutes filières/pays) – lecture directe de l'onglet EPISODES"
ws[f"A{r0 - 1}"].font = Font(bold=True)
for j, h in enumerate(["filiere", "pays", "sens", "nb_episodes", "beta_moyen", "choc_int_moyen", "choc_prod_moyen"], 1):
    ws.cell(row=r0, column=j, value=h).font = Font(bold=True)
seen = []
for r in D.EPISODES:
    k = (r["filiere"], r["pays"], r["sens"])
    if k not in seen:
        seen.append(k)
for i, (fil, pays, sens) in enumerate(seen, r0 + 1):
    ws[f"A{i}"], ws[f"B{i}"], ws[f"C{i}"] = fil, pays, sens
    crit = f'{erng("filiere")},A{i},{erng("pays")},B{i},{erng("sens")},C{i}'
    ws[f"D{i}"] = f"=COUNTIFS({crit})"
    ws[f"E{i}"] = f"=AVERAGEIFS({erng('transmission')},{crit})"
    ws[f"F{i}"] = f"=AVERAGEIFS({erng('choc_int')},{crit})"
    ws[f"G{i}"] = f"=AVERAGEIFS({erng('choc_prod')},{crit})"
    ws[f"E{i}"].number_format = "0.00"
    ws[f"F{i}"].number_format = ws[f"G{i}"].number_format = "0.0%"

# ---------------------------------------------------------------- STRESS
ws = wb.create_sheet("STRESS_PARAM")
header(ws, 1, D.STRESS_COLS + ["valeur_filiere_mds"], [16, 8, 14, 9, 7, 7, 10, 10, 50, 80, 12])
for i, r in enumerate(D.STRESS, 2):
    for j, k in enumerate(D.STRESS_COLS, 1):
        c = ws.cell(row=i, column=j, value=r[k])
        if k in ("p0", "q0_kt", "beta1", "beta2", "part_publique", "capacite_mds"):
            c.font = INPUT_FONT
    ws[f"K{i}"] = f"=B{i}*D{i}/1000"
    ws[f"K{i}"].number_format = "#,##0"
ws["A10"] = ("Hypothèses : beta1 = transmission pendant la campagne en cours ; beta2 = transmission à la campagne suivante ; "
             "part_publique = part de la charge d'amortissement susceptible de revenir à l'État/régulateur ; "
             "capacité = réserves mobilisables publiées (0 si non publiées).")

ws = wb.create_sheet("STRESS_TEST")
scols = ["filiere", "scenario", "libelle", "choc_int", "dQ", "deux_campagnes", "valeur_filiere_mds",
         "dP_prod_c1", "dP_prod_c2", "dRecette_c1", "perte_recette_c1_mds", "perte_recette_c2_mds", "perte_recette_cumulee_mds",
         "charge_absorbee_hors_prod_mds", "besoin_public_mds", "capacite_mds", "deficit_financement_mds",
         "besoin_public_FCFA_kg", "besoin_public_pct_valeur", "cout_protection_totale_c1_mds", "cout_garantie_recette_90_mds"]
header(ws, 1, scols, [16, 8, 30] + [12] * (len(scols) - 3))
row = 2
for pi, prm in enumerate(D.STRESS, 2):
    P = lambda col: f"STRESS_PARAM!${col}${pi}"  # noqa: E731
    for sc in D.SCENARIOS:
        r = row
        ws[f"A{r}"], ws[f"B{r}"], ws[f"C{r}"] = prm["filiere"], sc["code"], sc["nom"]
        ws[f"D{r}"], ws[f"E{r}"], ws[f"F{r}"] = sc["choc"], sc["dq"], 1 if sc["deux_campagnes"] else 0
        for c in "DEF":
            ws[f"{c}{r}"].font = INPUT_FONT
        ws[f"G{r}"] = f"={P('K')}"
        ws[f"H{r}"] = f"={P('E')}*D{r}"
        ws[f"I{r}"] = f'=IF(F{r}=1,{P("F")}*D{r},"")'
        ws[f"J{r}"] = f"=(1+H{r})*(1+E{r})-1"
        ws[f"K{r}"] = f"=-J{r}*G{r}"
        ws[f"L{r}"] = f"=IF(F{r}=1,-I{r}*G{r},0)"
        ws[f"M{r}"] = f"=K{r}+L{r}"
        ws[f"N{r}"] = f"=(1-{P('E')})*(-D{r})*G{r}*(1+E{r})+IF(F{r}=1,(1-{P('F')})*(-D{r})*G{r},0)"
        ws[f"O{r}"] = f"=N{r}*{P('G')}"
        ws[f"P{r}"] = f"={P('H')}"
        ws[f"Q{r}"] = f"=MAX(0,O{r}-P{r})"
        ws[f"R{r}"] = f"=O{r}*1000/({P('D')}*(1+E{r})*(1+F{r}))"
        ws[f"S{r}"] = f"=O{r}/G{r}"
        ws[f"T{r}"] = f"=-H{r}*G{r}*(1+E{r})"
        ws[f"U{r}"] = f"=MAX(0,0.9-(1+J{r}))*G{r}"
        for c in "DEHIJS":
            ws[f"{c}{r}"].number_format = "0.0%"
        for c in "GKLMNOPQRTU":
            ws[f"{c}{r}"].number_format = "#,##0"
        row += 1
ws.add_table(Table(displayName="T_STRESS", ref=f"A1:{L(len(scols))}{row - 1}",
                   tableStyleInfo=TableStyleInfo(name="TableStyleLight9", showRowStripes=True)))

# ---------------------------------------------------------------- MECANISMES / LITTERATURE / CARTE
for name, cols_, rows_, widths in [
    ("MECANISMES", D.MECA_COLS, D.MECANISMES, [12, 12, 8, 50, 40, 45, 50, 18, 8, 25]),
    ("LITTERATURE", D.LIT_COLS, D.LITTERATURE, [55, 22, 18, 55, 35, 45]),
    ("CARTE_QUALI", ["filiere", "exposition", "transmission", "absorbeur", "vulnerabilite"], D.FILIERES_QUALI, [16, 28, 50, 40, 55]),
]:
    ws = wb.create_sheet(name)
    header(ws, 1, cols_, widths)
    for i, r in enumerate(rows_, 2):
        for j, v in enumerate(r, 1):
            ws.cell(row=i, column=j, value=v).alignment = Alignment(wrap_text=True, vertical="top")

wb.save(OUT)
print("écrit", OUT)
