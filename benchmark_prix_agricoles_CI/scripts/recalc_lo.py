"""Recalcule toutes les formules du classeur avec LibreOffice sans interface, puis réenregistre le .xlsx.

openpyxl écrit les formules sans valeurs calculées : ce recalcul est nécessaire avant de lire les
résultats (make_charts.py, make_note.py, export_csv.py) ou d'ouvrir le fichier dans un lecteur
qui n'évalue pas les formules.

Usage     : python3 scripts/recalc_lo.py [Base_Prix_Benchmark_CI.xlsx]
Prérequis : LibreOffice (soffice) et son pont Python « uno » (paquet python3-uno sous Debian/Ubuntu).
Alternative manuelle : ouvrir le classeur dans Excel ou LibreOffice, forcer le recalcul
(Ctrl+Alt+F9 / Ctrl+Maj+F9) et enregistrer au format .xlsx.
"""
import os
import shutil
import subprocess
import sys
import tempfile
import time
from pathlib import Path

import uno
from com.sun.star.beans import PropertyValue
from openpyxl import load_workbook

ERREURS = ("#VALUE!", "#DIV/0!", "#REF!", "#NAME?", "#N/A", "#NUM!", "#NULL!")


def prop(name, value):
    p = PropertyValue()
    p.Name, p.Value = name, value
    return p


def recalc(path: Path) -> None:
    profile = Path(tempfile.mkdtemp(prefix="lo_profil_"))
    port = 20000 + os.getpid() % 10000
    env = dict(os.environ, SAL_USE_VCLPLUGIN="svp")
    proc = subprocess.Popen(
        ["soffice", "--headless", "--invisible", "--nologo", "--norestore", "--nodefault",
         f"-env:UserInstallation={profile.as_uri()}",
         f"--accept=socket,host=127.0.0.1,port={port};urp;"],
        stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, env=env)
    desktop = None
    try:
        local = uno.getComponentContext()
        resolver = local.ServiceManager.createInstanceWithContext("com.sun.star.bridge.UnoUrlResolver", local)
        for _ in range(120):
            try:
                ctx = resolver.resolve(f"uno:socket,host=127.0.0.1,port={port};urp;StarOffice.ComponentContext")
                break
            except Exception:
                if proc.poll() is not None:
                    raise RuntimeError("LibreOffice s'est arrêté au démarrage")
                time.sleep(0.5)
        else:
            raise RuntimeError("LibreOffice ne répond pas")
        desktop = ctx.ServiceManager.createInstanceWithContext("com.sun.star.frame.Desktop", ctx)
        doc = desktop.loadComponentFromURL(path.as_uri(), "_blank", 0, (prop("Hidden", True),))
        doc.calculateAll()
        doc.storeToURL(path.as_uri(), (prop("FilterName", "Calc MS Excel 2007 XML"), prop("Overwrite", True)))
        doc.close(True)
    finally:
        try:
            if desktop is not None:
                desktop.terminate()
        except Exception:
            pass
        try:
            proc.wait(timeout=30)
        except subprocess.TimeoutExpired:
            proc.kill()
        shutil.rmtree(profile, ignore_errors=True)


def controle(path: Path) -> int:
    formules = load_workbook(path)
    valeurs = load_workbook(path, data_only=True)
    n_form, erreurs = 0, []
    for ws in formules.worksheets:
        wv = valeurs[ws.title]
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    n_form += 1
                    v = wv[c.coordinate].value
                    if isinstance(v, str) and v in ERREURS:
                        erreurs.append(f"{ws.title}!{c.coordinate} {v}")
    print(f"{n_form} formules recalculées ; {len(erreurs)} erreur(s)")
    for e in erreurs[:50]:
        print("  ", e)
    return len(erreurs)


if __name__ == "__main__":
    cible = Path(sys.argv[1] if len(sys.argv) > 1 else Path(__file__).resolve().parent.parent / "Base_Prix_Benchmark_CI.xlsx").resolve()
    recalc(cible)
    sys.exit(1 if controle(cible) else 0)
