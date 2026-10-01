"""Exports CSV de la base recalculée (valeurs) et liste des sources.

Entrée  : Base_Prix_Benchmark_CI.xlsx (déjà recalculé : valeurs en cache).
Sorties : Base_Prix_Benchmark_CI_observations.csv
          SOURCES/Liste_sources.csv
          bloc « Liste des sources » de METHODOLOGIE_ET_SOURCES.md (entre les balises SOURCES:DEBUT / SOURCES:FIN)

Format CSV : séparateur « ; », décimale « , », UTF-8 avec BOM (ouverture directe dans Excel en français).
Lecture Python : pd.read_csv(fichier, sep=";", decimal=",")
"""
import csv
import re
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
wb = load_workbook(ROOT / "Base_Prix_Benchmark_CI.xlsx", data_only=True)


def fmt(v):
    if v is None:
        return ""
    if isinstance(v, float):
        return f"{v:.6g}".replace(".", ",") if abs(v) < 1e15 else str(v)
    return str(v)


def rows_of(ws):
    hdr = [c.value for c in ws[1]]
    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is None:
            continue
        out.append(list(r[: len(hdr)]))
    return hdr, out


# --- Observations ----------------------------------------------------------
hdr, obs = rows_of(wb["Observations"])
with open(ROOT / "Base_Prix_Benchmark_CI_observations.csv", "w", newline="", encoding="utf-8-sig") as f:
    w = csv.writer(f, delimiter=";")
    w.writerow(hdr)
    for r in obs:
        w.writerow([fmt(v) for v in r])

# --- Sources ---------------------------------------------------------------
# Sources utilisées hors de l'onglet Observations : emplacement décrit à la main (taux de change, décomposition,
# chiffres de contexte cités dans la note, recoupements).
CONTEXTE = {
    "BCE_FX": "Change_mensuel, Change_periodes (taux BCE : EUR donc FCFA ; IDR, THB, MYR)",
    "BRI_FX": "Change_mensuel, Change_periodes (GHS, UGX, VND, TZS, NGN)",
    "BM_CACAO_2019": "Decomposition (règle de partage du CAF, prélèvements ≈22 %, marges privées)",
    "OMC_TPR_2017": "Decomposition (DUS 14,6 % du CAF ; prélèvements 2016/17)",
    "CI_ANA_DUS": "Decomposition (DUS de la noix brute : 5 %)",
    "CI_CAO_MECA": "Decomposition (part producteur caoutchouc : 63 % puis 66 %)",
    "CI_CAC_STOCKS": "Note p. 3 et 5 (123 000 t invendues en 2025/26)",
    "CI_CAC_VENTES": "Note p. 3 et 4 ; Matrice ; Cacao_transmission (ventes anticipées mars-juin 2026)",
    "GH_CAC_GAP": "Note p. 3 et 5 (déficit de financement du COCOBOD ≈1,4 Md $)",
    "GH_CAC_REFORM": "Contexte Ghana (réformes du COCOBOD), consulté pour l'analyse",
    "ECOFIN_GAP": "Recoupement de l'écart de prix Ghana / Côte d'Ivoire 2026/27",
    "CI_ANA_EXP": "Contexte anacarde (exportations 2025), consulté pour l'analyse",
    "CI_ANA_TRANSFO": "Contexte anacarde (achats réservés aux transformateurs locaux), consulté pour l'analyse",
    "CI_ANN_EXP": "Note p. 4 (exportations d'ananas 2023)",
    "CI_BAN_EXP": "Note p. 4 (exportations de bananes 2025)",
    "OCPV": "Note p. 4 ; Matrice (prix à la consommation : stade non comparable)",
    "USDA_PAL": "Contexte palmier (production d'huile), consulté pour l'analyse",
    "USDA_RIZ_IMP": "Note p. 4 (importations de riz ≈1,75 Mt)",
}
i_src = hdr.index("Source")
shdr, src = rows_of(wb["Sources"])
usage = Counter()
cles = [s[0] for s in src]
inst_titre = {s[0]: f"{s[1]} – {s[2]}" for s in src}
for r in obs:
    lib = r[i_src] or ""
    for k, lt in inst_titre.items():
        if lib == lt:
            usage[k] += 1
textes = {}
for ws in load_workbook(ROOT / "Base_Prix_Benchmark_CI.xlsx").worksheets:  # valeurs et commentaires de cellule
    if ws.title in ("Sources", "Lisez-moi"):
        continue
    buf = []
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None:
                buf.append(str(c.value))
            if c.comment:
                buf.append(c.comment.text)
    textes[ws.title] = "\n".join(buf)
emploi = {}
for s in src:
    k, url = s[0], s[3] or ""
    onglets = [t for t, x in textes.items() if (url.startswith("http") and url in x) or inst_titre[k] in x]
    emploi[k] = " ; ".join(x for x in [", ".join(onglets), CONTEXTE.get(k, "")] if x) or "—"
with open(ROOT / "SOURCES" / "Liste_sources.csv", "w", newline="", encoding="utf-8-sig") as f:
    w = csv.writer(f, delimiter=";")
    w.writerow(shdr + ["Nb d'observations (onglet Observations)", "Utilisation (onglets de la base / note)"])
    for s in src:
        w.writerow([fmt(v) for v in s] + [usage.get(s[0], 0), emploi[s[0]]])

# --- Bloc Markdown ---------------------------------------------------------
md_path = ROOT / "METHODOLOGIE_ET_SOURCES.md"
if md_path.exists():
    def cellmd(t):
        return str(t or "").replace("|", "/").replace("\n", " ")

    lines = [
        f"{len(src)} sources (fiabilité A : {sum(1 for s in src if s[5] == 'A')} ; "
        f"B : {sum(1 for s in src if s[5] == 'B')} ; C : {sum(1 for s in src if s[5] == 'C')}), "
        "toutes consultées le 1er octobre 2026 (date indiquée pour chaque ligne de la base). "
        "Version tableur : `SOURCES/Liste_sources.csv` (avec le nombre d'observations qui citent chaque source).",
        "",
        "| Clé | Institution | Titre / contenu utilisé | Fiab. | Utilisation | Lien |",
        "|---|---|---|---|---|---|",
    ]
    for s in src:
        url = s[3] or ""
        lien = f"[lien]({url})" if url.startswith("http") else cellmd(url)
        n = usage.get(s[0], 0)
        util = f"{n} obs." + (f" ; {emploi[s[0]]}" if emploi[s[0]] not in ("Observations", "—") else "") if n else emploi[s[0]]
        lines.append(f"| `{s[0]}` | {cellmd(s[1])} | {cellmd(s[2])} | {s[5]} | {cellmd(util)} | {lien} |")
    block = "\n".join(lines)
    md = md_path.read_text(encoding="utf-8")
    md = re.sub(r"(<!-- SOURCES:DEBUT -->\n).*?(<!-- SOURCES:FIN -->)", lambda m: m.group(1) + block + "\n" + m.group(2),
                md, flags=re.S)
    md_path.write_text(md, encoding="utf-8")

print(f"{len(obs)} observations, {len(src)} sources exportées ; sources sans emplacement identifié :",
      [k for k in cles if emploi[k] == "—"])
